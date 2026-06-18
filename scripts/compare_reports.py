from docx import Document
import openpyxl

# 1. Read Test Plan v1.2 test case index (Table 8)
plan = Document(r"D:\Downloads\Test Plan v1.2.docx")
t8 = plan.tables[7]
plan_cases = []
for r_idx, row in enumerate(t8.rows):
    if r_idx == 0:
        continue
    cells = [c.text.strip() for c in row.cells]
    plan_cases.append({
        "group": cells[0],
        "jira": cells[1],
        "tc_id": cells[2],
        "scenario": cells[3],
        "level": cells[4],
        "report_ref": cells[5],
    })

print("=== TEST PLAN v1.2: " + str(len(plan_cases)) + " test cases ===")
for c in plan_cases:
    print("  " + c["group"] + " | " + c["tc_id"] + " | " + c["scenario"][:50] + " | " + c["report_ref"])

# 2. Read Test Report Excel
print()
print("=== TEST REPORT EXCEL ===")
wb = openpyxl.load_workbook(r"D:\Downloads\Test Report.xlsx")
ws = wb["Test Cases"]
print("Rows: " + str(ws.max_row) + ", Cols: " + str(ws.max_column))
# Headers
headers = [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
print("Headers: " + " | ".join(headers))
print()
# All data rows
excel_cases = []
for r in range(2, ws.max_row + 1):
    case_num = str(ws.cell(r, 1).value or "")
    scenario = str(ws.cell(r, 2).value or "")[:50]
    status = str(ws.cell(r, ws.max_column).value or "")
    excel_cases.append(case_num)
    print("  " + case_num + " | " + scenario + " | " + status)

# 3. Read Test Report Word doc
print()
print("=== TEST REPORT WORD ===")
try:
    report = Document(r"D:\Downloads\Test Report - ISR updated.docx")
    # Find all tables and look for test case tables (4 cols with TC ID)
    word_tc_ids = []
    for t_idx, table in enumerate(report.tables):
        if len(table.columns) == 4 and len(table.rows) > 1:
            h = [c.text.strip() for c in table.rows[0].cells]
            if "TC ID" in h[0] or "TC ID" in str(h):
                for r_idx in range(1, len(table.rows)):
                    tc_id = table.rows[r_idx].cells[0].text.strip()
                    feature = table.rows[r_idx].cells[1].text.strip()[:40]
                    status = table.rows[r_idx].cells[3].text.strip()
                    word_tc_ids.append(tc_id)
                    print("  " + tc_id + " | " + feature + " | " + status)
    print("Total TCs in Word report: " + str(len(word_tc_ids)))
except Exception as e:
    print("Could not read Word report: " + str(e))

# 4. Compare
print()
print("=== COMPARISON ===")
plan_tc_ids = set()
for c in plan_cases:
    for tc in c["tc_id"].replace("\n", ",").split(","):
        plan_tc_ids.add(tc.strip())
print("Plan TC IDs: " + str(len(plan_tc_ids)))
print("Excel cases: " + str(len(excel_cases)))

# Check plan report_ref vs excel case numbers
plan_refs = set()
for c in plan_cases:
    for ref in c["report_ref"].replace("\n", ",").split(","):
        plan_refs.add(ref.strip())
print("Plan report refs: " + str(sorted(plan_refs)[:10]) + "...")
print("Excel case nums: " + str(excel_cases[:10]) + "...")
