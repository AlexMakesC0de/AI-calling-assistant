import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
import openpyxl

# Plan
plan = Document(r"D:\Downloads\Test Plan v1.2.docx")
t8 = plan.tables[7]
plan_refs = []
for r_idx, row in enumerate(t8.rows):
    if r_idx == 0:
        continue
    cells = [c.text.strip() for c in row.cells]
    refs = cells[5].replace("\n", ",").split(",")
    for ref in refs:
        ref = ref.strip().replace(" ", "").replace("# ", "#")
        if ref:
            plan_refs.append(ref)

# Excel
wb = openpyxl.load_workbook(r"D:\Downloads\Test Report.xlsx")
ws = wb["Test Cases"]
excel_refs = []
for r in range(2, ws.max_row + 1):
    val = str(ws.cell(r, 1).value or "").strip()
    if val:
        excel_refs.append(val)

print("Plan refs (" + str(len(plan_refs)) + "):")
for r in plan_refs:
    print("  " + r)

print()
print("Excel refs (" + str(len(excel_refs)) + "):")
for r in excel_refs:
    print("  " + r)

print()
plan_set = set(plan_refs)
excel_set = set(excel_refs)
print("In Plan not in Excel: " + str(sorted(plan_set - excel_set)))
print("In Excel not in Plan: " + str(sorted(excel_set - plan_set)))
print("Match: " + str(plan_set == excel_set))

# Check order matches
print()
print("Order match: " + str(plan_refs == excel_refs))
if plan_refs != excel_refs:
    for i in range(min(len(plan_refs), len(excel_refs))):
        if plan_refs[i] != excel_refs[i]:
            print("First mismatch at index " + str(i) + ": plan=" + plan_refs[i] + " excel=" + excel_refs[i])
            break
