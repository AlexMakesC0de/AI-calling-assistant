import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
import openpyxl

# 1. Test Plan v1.2 test case index
plan = Document(r"D:\Downloads\Test Plan v1.2.docx")
t8 = plan.tables[7]
plan_cases = []
for r_idx, row in enumerate(t8.rows):
    if r_idx == 0:
        continue
    cells = [c.text.strip() for c in row.cells]
    # Parse TC IDs (some rows have multiple like "TC 22.1,\nTC 22.2, TC 22.3")
    tc_ids_raw = cells[2].replace("\n", ",")
    for tc in tc_ids_raw.split(","):
        tc = tc.strip()
        if tc:
            plan_cases.append(tc)

# Also collect report references
plan_refs = []
for r_idx, row in enumerate(t8.rows):
    if r_idx == 0:
        continue
    cells = [c.text.strip() for c in row.cells]
    refs_raw = cells[5].replace("\n", ",")
    for ref in refs_raw.split(","):
        ref = ref.strip()
        if ref:
            plan_refs.append(ref)

# 2. Excel report
wb = openpyxl.load_workbook(r"D:\Downloads\Test Report.xlsx")
ws = wb["Test Cases"]
excel_cases = []
excel_scenarios = []
for r in range(2, ws.max_row + 1):
    case_num = str(ws.cell(r, 1).value or "").strip()
    scenario = str(ws.cell(r, 2).value or "").strip()
    status = str(ws.cell(r, 10).value or "").strip()
    if case_num:
        excel_cases.append(case_num)
        excel_scenarios.append((case_num, scenario[:60], status))

# 3. Word report
report = Document(r"D:\Downloads\Test Report - ISR updated.docx")
word_cases = []
for t in report.tables:
    if len(t.columns) == 4 and len(t.rows) > 1:
        h0 = t.rows[0].cells[0].text.strip()
        if "TC" in h0 or "TC" in t.rows[0].cells[0].text:
            for r_idx in range(1, len(t.rows)):
                tc_id = t.rows[r_idx].cells[0].text.strip()
                word_cases.append(tc_id)

# 4. Comparison
print("=== COUNTS ===")
print("Test Plan v1.2 TC IDs: " + str(len(plan_cases)))
print("Test Plan v1.2 report refs: " + str(len(plan_refs)))
print("Excel rows: " + str(len(excel_cases)))
print("Word report TCs: " + str(len(word_cases)))

# Normalize plan refs to match excel format
# Plan uses "Sprint 2 #1", Excel uses "Sprint2#1"
def normalize_ref(ref):
    return ref.replace(" ", "").replace("#", "#")

plan_refs_norm = set(normalize_ref(r) for r in plan_refs)
excel_refs = set(excel_cases)

print()
print("=== IN PLAN BUT NOT IN EXCEL ===")
missing_from_excel = sorted(plan_refs_norm - excel_refs)
for m in missing_from_excel:
    print("  " + m)
print("Count: " + str(len(missing_from_excel)))

print()
print("=== IN EXCEL BUT NOT IN PLAN ===")
extra_in_excel = sorted(excel_refs - plan_refs_norm)
for e in extra_in_excel:
    print("  " + e)
print("Count: " + str(len(extra_in_excel)))

print()
print("=== PLAN TC IDs vs WORD REPORT TC IDs ===")
plan_tc_set = set(plan_cases)
word_tc_set = set(word_cases)
missing_from_word = sorted(plan_tc_set - word_tc_set)
extra_in_word = sorted(word_tc_set - plan_tc_set)
print("In Plan but not in Word: " + str(len(missing_from_word)))
for m in missing_from_word[:10]:
    print("  " + m)
print("In Word but not in Plan: " + str(len(extra_in_word)))
for e in extra_in_word[:10]:
    print("  " + e)

# Show Excel sprint5/6 rows
print()
print("=== EXCEL SPRINT 5-6 ROWS ===")
for case_num, scenario, status in excel_scenarios:
    if "Sprint5" in case_num or "Sprint6" in case_num or "5#" in case_num or "6#" in case_num:
        print("  " + case_num + " | " + scenario + " | " + status)
