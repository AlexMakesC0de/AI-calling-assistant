"""Tests for non-image attachment references in the .docx (ISR-319)."""

from __future__ import annotations

import pytest
from docx import Document

import generate_word as gw


def _render(tmp_path, attachments):
    doc = Document()
    gw.embed_attachments(doc, {"attachments": attachments}, str(tmp_path))
    out = tmp_path / "out.docx"
    doc.save(str(out))
    return Document(str(out))


def _all_text(doc):
    return "\n".join(p.text for p in doc.paragraphs)


# --- type label helper -----------------------------------------------------

@pytest.mark.parametrize("name,label", [
    ("report.pdf", "PDF"),
    ("notes.docx", "DOCX"),
    ("sheet.xlsx", "XLSX"),
    ("archive.zip", "ZIP"),
    ("clip.mp4", "MP4"),
    ("data.xyz", "XYZ"),     # unknown extension -> upper-cased ext
    ("README", "FILE"),      # no extension -> FILE
])
def test_file_type_label(name, label):
    assert gw._file_type_label(name) == label


# --- rendering -------------------------------------------------------------

def test_pdf_reference_lists_filename_type_and_path(tmp_path):
    out = _render(tmp_path, [{"path": "docs/crash.pdf", "filename": "crash-report.pdf"}])
    text = _all_text(out)
    assert "crash-report.pdf" in text          # filename
    assert "[PDF]" in text                      # type clearly indicated
    assert "crash.pdf" in text and "docs" in text  # path present & locally accessible


def test_non_image_is_not_embedded_as_picture(tmp_path):
    (tmp_path / "report.pdf").write_bytes(b"%PDF-1.4 fake")
    out = _render(tmp_path, ["report.pdf"])
    assert len(out.inline_shapes) == 0
    assert "[PDF]" in _all_text(out)


def test_filename_is_bold(tmp_path):
    out = _render(tmp_path, [{"path": "x/notes.docx", "filename": "notes.docx"}])
    bold_texts = [
        r.text for p in out.paragraphs for r in p.runs if r.bold
    ]
    assert any("notes.docx" in t for t in bold_texts)


def test_docx_type_label_rendered(tmp_path):
    out = _render(tmp_path, ["meeting-notes.docx"])
    assert "[DOCX]" in _all_text(out)


def test_path_is_resolved_against_media_dir(tmp_path):
    out = _render(tmp_path, ["sub/manual.pdf"])
    # the rendered path should point at the resolved, locally-accessible file
    assert str(tmp_path / "sub" / "manual.pdf") in _all_text(out)


def test_mixed_image_and_file_attachments(tmp_path):
    # one real image + one pdf reference
    import struct
    import zlib

    def _png(path, w, h):
        def chunk(typ, data):
            body = typ + data
            return struct.pack(">I", len(data)) + body + struct.pack(
                ">I", zlib.crc32(body) & 0xFFFFFFFF
            )
        ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)
        raw = (b"\x00" + b"\xff\x00\x00" * w) * h
        path.write_bytes(
            b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", zlib.compress(raw))
            + chunk(b"IEND", b"")
        )

    _png(tmp_path / "shot.png", 200, 200)
    out = _render(tmp_path, ["shot.png", {"path": "r.pdf", "filename": "r.pdf"}])
    assert len(out.inline_shapes) == 1          # only the image embedded
    assert "[PDF]" in _all_text(out)            # pdf referenced
