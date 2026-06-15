"""Tests for image/attachment embedding in the .docx generator (ISR-318)."""

from __future__ import annotations

import struct
import zlib

import pytest
from docx import Document

import generate_word as gw


# --- helpers ---------------------------------------------------------------

def _make_png(path, width, height, color=(200, 30, 30)):
    """Write a minimal valid RGB PNG of the given pixel size (no PIL needed)."""
    def chunk(typ, data):
        body = typ + data
        return (
            struct.pack(">I", len(data))
            + body
            + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)  # 8-bit RGB
    row = b"\x00" + bytes(color) * width
    raw = row * height
    png = (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )
    path.write_bytes(png)
    return str(path)


def _doc_with_token(tmp_path, token_text=None):
    """A tiny template; optionally include the {attachments} token paragraph."""
    doc = Document()
    doc.add_paragraph("Incident form")
    if token_text is not None:
        doc.add_paragraph(token_text)
    p = tmp_path / "tmpl.docx"
    doc.save(str(p))
    return str(p)


def _render(tmp_path, data, template=None, media_dir=None):
    """Run the full token + attachment pipeline and return the saved Document."""
    doc = Document(template) if template else Document()
    gw.replace_tokens_in_doc(doc, data)
    gw.embed_attachments(doc, data, media_dir or str(tmp_path))
    out = tmp_path / "out.docx"
    doc.save(str(out))
    return Document(str(out))


# --- _collect_attachments --------------------------------------------------

def test_collect_handles_string_and_dict_entries(tmp_path):
    data = {"attachments": [
        "photo.png",
        {"path": "scan.jpg", "filename": "Scan 1.jpg"},
        {"filename": "report.pdf"},
    ]}
    items = gw._collect_attachments(data, str(tmp_path))
    assert len(items) == 3
    assert items[0][2] is True   # png -> image
    assert items[1][1] == "Scan 1.jpg"
    assert items[2][2] is False  # pdf -> not image


def test_collect_resolves_relative_paths(tmp_path):
    items = gw._collect_attachments({"attachments": ["a.png"]}, str(tmp_path))
    assert items[0][0] == str(tmp_path / "a.png")


def test_collect_empty_when_no_attachments():
    assert gw._collect_attachments({}, None) == []


# --- embed_attachments: images ---------------------------------------------

def test_embeds_multiple_images(tmp_path):
    _make_png(tmp_path / "a.png", 400, 200)
    _make_png(tmp_path / "b.png", 400, 200)
    out = _render(tmp_path, {"attachments": ["a.png", "b.png"]})
    assert len(out.inline_shapes) == 2  # both images embedded


def test_no_attachments_adds_no_pictures(tmp_path):
    out = _render(tmp_path, {"attachments": []})
    assert len(out.inline_shapes) == 0


def test_large_image_scaled_to_page_width(tmp_path):
    _make_png(tmp_path / "wide.png", 4000, 1000)  # very wide
    doc = Document()
    usable = gw._usable_width(doc)
    out = _render(tmp_path, {"attachments": ["wide.png"]})
    shape = out.inline_shapes[0]
    assert shape.width <= usable
    # aspect ratio preserved (4:1 within rounding)
    assert abs((shape.width / shape.height) - 4.0) < 0.05


def test_small_image_not_upscaled(tmp_path):
    _make_png(tmp_path / "small.png", 60, 40)
    doc = Document()
    usable = gw._usable_width(doc)
    out = _render(tmp_path, {"attachments": ["small.png"]})
    shape = out.inline_shapes[0]
    assert shape.width < usable  # left at (small) native size, not blown up


# --- anchor token vs append ------------------------------------------------

def test_attachments_token_is_consumed_and_used_as_anchor(tmp_path):
    _make_png(tmp_path / "a.png", 300, 300)
    template = _doc_with_token(tmp_path, "{attachments}")
    out = _render(tmp_path, {"attachments": ["a.png"]}, template=template)
    # the literal token must be gone, replaced by the Attachments heading
    texts = [p.text for p in out.paragraphs]
    assert "{attachments}" not in " ".join(texts)
    assert any("Attachments" in t for t in texts)
    assert len(out.inline_shapes) == 1


def test_appends_section_when_no_token(tmp_path):
    _make_png(tmp_path / "a.png", 300, 300)
    out = _render(tmp_path, {"attachments": ["a.png"]})
    assert any("Attachments" in p.text for p in out.paragraphs)
    assert len(out.inline_shapes) == 1


# --- robustness ------------------------------------------------------------

def test_missing_image_does_not_crash(tmp_path):
    out = _render(tmp_path, {"attachments": ["nope.png"]})
    assert len(out.inline_shapes) == 0
    assert any("Missing image" in p.text for p in out.paragraphs)


def test_non_image_attachment_referenced_not_embedded(tmp_path):
    (tmp_path / "report.pdf").write_bytes(b"%PDF-1.4 fake")
    out = _render(tmp_path, {"attachments": ["report.pdf"]})
    assert len(out.inline_shapes) == 0
    assert any("report.pdf" in p.text for p in out.paragraphs)


def test_corrupt_image_does_not_break_document(tmp_path):
    (tmp_path / "broken.png").write_bytes(b"not a real png")
    out = _render(tmp_path, {"attachments": ["broken.png"]})
    # no picture, but the doc still saved/opened and noted the failure
    assert len(out.inline_shapes) == 0
    assert any("Could not embed image" in p.text for p in out.paragraphs)


def test_tokens_still_replaced_alongside_attachments(tmp_path):
    _make_png(tmp_path / "a.png", 200, 200)
    doc = Document()
    doc.add_paragraph("Caller: {caller_information.name}")
    template = tmp_path / "t.docx"
    doc.save(str(template))
    out = _render(
        tmp_path,
        {"caller_information": {"name": "Jane Doe"}, "attachments": ["a.png"]},
        template=str(template),
    )
    assert any("Jane Doe" in p.text for p in out.paragraphs)
    assert len(out.inline_shapes) == 1
