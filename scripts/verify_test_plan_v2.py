from docx import Document

doc = Document(r"D:\Downloads\Test Plan v1.2.docx")

print("=== VERIFICATION ===")
print("Total paragraphs: " + str(len(doc.paragraphs)))
print("Total tables: " + str(len(doc.tables)))
print()

# 1. Version History
t1 = doc.tables[0]
print("--- Version History (Table 1, " + str(len(t1.rows)) + " rows) ---")
last_row = t1.rows[-1]
cells = [c.text.strip()[:50] for c in last_row.cells]
print("  Last row: " + " | ".join(cells))
print()

# 2. Test Items
t2 = doc.tables[1]
print("--- Test Items (Table 2, " + str(len(t2.rows)) + " rows) ---")
for r in t2.rows[-3:]:
    cells = [c.text.strip()[:40] for c in r.cells]
    print("  " + " | ".join(cells))
print()

# 3. Requirements
t3 = doc.tables[2]
print("--- Requirements (Table 3, " + str(len(t3.rows)) + " rows) ---")
for r in t3.rows[-3:]:
    cells = [c.text.strip()[:50] for c in r.cells]
    print("  " + " | ".join(cells))
print()

# 4. Test Case Index
t8 = doc.tables[7]
print("--- Test Case Index (Table 8, " + str(len(t8.rows)) + " rows) ---")
print("  First new row:")
if len(t8.rows) > 25:
    cells = [c.text.strip()[:50] for c in t8.rows[25].cells]
    print("  " + " | ".join(cells))
print("  Last 3 rows:")
for r in t8.rows[-3:]:
    cells = [c.text.strip()[:50] for c in r.cells]
    print("  " + " | ".join(cells))
print()

# 5. Check new spec tables exist
print("--- New Spec Tables ---")
# The original had tables 9-35 (indices 8-34) = 27 spec tables
# New ones start after that
# Count spec-like tables (2 cols, 7 rows)
spec_count = 0
for t in doc.tables:
    if len(t.columns) == 2 and len(t.rows) == 7:
        first_cell = t.rows[0].cells[0].text.strip()
        if first_cell.startswith("Technique"):
            spec_count += 1
print("  Total spec card tables (2x7 with 'Technique'): " + str(spec_count))
print()

# 6. Check headings for new spec cards
print("--- New Heading 3s (last 5) ---")
h3s = []
for p in doc.paragraphs:
    if p.style.name == "Heading 3":
        h3s.append(p.text[:80])
for h in h3s[-5:]:
    print("  " + h)
print("  Total Heading 3s: " + str(len(h3s)))
print()

# 7. Risks
print("--- Looking for risk tables ---")
for t_idx, t in enumerate(doc.tables):
    if len(t.rows) > 0:
        first = t.rows[0].cells[0].text.strip()[:20]
        if first in ["Level", "Risk"]:
            print("  Table " + str(t_idx + 1) + " (" + str(len(t.rows)) + " rows x " + str(len(t.columns)) + " cols): " + first)
            if len(t.rows) > 1:
                last = [c.text.strip()[:30] for c in t.rows[-1].cells]
                print("    Last row: " + " | ".join(last))
