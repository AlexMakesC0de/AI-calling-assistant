"""
modify_test_plan.py

Opens the existing Test Plan v1.1.docx and appends Sprint 5-6 content
(test items, requirements, test case index, and risks) without modifying
any existing content.  Saves as Test Plan v1.3.docx.
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT

INPUT_PATH = r"D:\Downloads\Test Plan v1.1.docx"
OUTPUT_PATH = r"D:\Downloads\Test Plan v1.3.docx"


def set_cell_text(cell, text, bold=False):
    """Set cell text, clearing default empty paragraph."""
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold


def add_table(doc, headers, rows):
    """Add a formatted table to the document."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Header row
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)

    # Data rows
    for r_idx, row_data in enumerate(rows):
        for c_idx, value in enumerate(row_data):
            set_cell_text(table.rows[r_idx + 1].cells[c_idx], value)

    return table


def main():
    doc = Document(INPUT_PATH)

    # ── 1. Version History Addition ──────────────────────────────────────
    doc.add_paragraph("")  # spacer
    doc.add_paragraph(
        "Version 1.3 — 14.06.2026 — Added test items, requirements, and test cases "
        "for Sprint 5-6 features: Outlook Email pipeline, junk filter, attachments, "
        "WhatsApp integration, Phone Calls, AI Output Reliability, Word Generator, "
        "and Dashboard."
    )

    # ── 2. New Test Items ────────────────────────────────────────────────
    doc.add_heading("Addendum: New Test Items (Sprint 5-6)", level=1)

    add_table(
        doc,
        ["Test Item", "Description", "Risks"],
        [
            [
                "Email Classifier",
                "Ollama LLM-based classification of emails as support/not-support/unclear",
                "False positives create junk cases; threshold sensitivity",
            ],
            [
                "WhatsApp Ingest",
                "Twilio webhook receiver for WhatsApp text, voice notes, and media",
                "Signature validation failure; media download timeout",
            ],
            [
                "Dashboard Auth",
                "Session-based login gating all dashboard routes",
                "Session bypass exposes customer data",
            ],
            [
                "Semantic Search",
                "pgvector cosine similarity search over transcript embeddings",
                "Embedding quality dependent on Ollama model",
            ],
            [
                "Email Signature Stripper",
                "Preprocessing to strip signatures and quoted replies from emails",
                "Over-stripping removes relevant content",
            ],
            [
                "Gemma Model",
                "Alternative AI model benchmarking against llama3.1",
                "Quality regression on edge cases",
            ],
        ],
    )

    # ── 3. New Requirements ──────────────────────────────────────────────
    doc.add_heading("Addendum: New Requirements (US-25 to US-40)", level=1)

    add_table(
        doc,
        ["ID", "Type", "Description"],
        [
            ["US-25", "Should-Have", "Auto-reply sent with case reference on inbound support email"],
            ["US-26", "Should-Have", "Inbound emails classified as support/not-support/unclear before case creation"],
            ["US-27", "Should-Have", "Email attachments downloaded and stored per case"],
            ["US-28", "Must-Have", "WhatsApp messages received and stored via Twilio webhook"],
            ["US-29", "Should-Have", "First message from new number triggers opt-in auto-reply"],
            ["US-30", "Must-Have", "Real provider calls captured and processed without changing existing number"],
            ["US-31", "Must-Have", "Follow-up interactions auto-linked to the same case"],
            ["US-32", "Must-Have", "Each issue has one growing service form"],
            ["US-33", "Should-Have", "Engineers can reassign, merge, or split cases"],
            ["US-34", "Must-Have", "Service form uses Repak's Problem/Analysis/Advice format"],
            ["US-35", "Should-Have", "Confirmation email sent after each interaction"],
            ["US-36", "Should-Have", "Bad AI response handled with retry and placeholder"],
            ["US-37", "Should-Have", "Dashboard login gates all routes"],
            ["US-38", "Should-Have", "Engineers can browse and review cases"],
            ["US-39", "Should-Have", "Hardware spec document for on-prem install"],
            ["US-40", "Must-Have", "Inbound support emails flow through the pipeline"],
        ],
    )

    # ── 4. Test Case Index ───────────────────────────────────────────────
    doc.add_heading("Addendum: Test Case Index (Sprint 5-6)", level=1)

    add_table(
        doc,
        ["#", "Jira Issue", "Test Case ID", "Test Scenario", "Test Level", "Epic"],
        [
            ["28", "ISR-238", "TC 25.1", "Email Pipeline E2E", "Integration", "Outlook Email"],
            ["29", "ISR-243", "TC 25.2", "Email Body to Formatter", "Component", "Outlook Email"],
            ["30", "ISR-311", "TC 26.1", "Email Classification", "Component", "Outlook Email"],
            ["31", "ISR-311", "TC 26.2", "Low-Confidence Email Flagging", "Integration", "Outlook Email"],
            ["32", "ISR-317", "TC 27.1", "Email Attachments Stored", "Component", "Outlook Email"],
            ["33", "ISR-320", "TC 27.2", "Oversized Attachment Handling", "Integration", "Outlook Email"],
            ["34", "ISR-346", "TC 25.6", "Email Signature Stripping", "Component", "Outlook Email"],
            ["35", "ISR-349", "TC 25.7", "MIME Email Support", "Component", "Outlook Email"],
            ["36", "ISR-226", "TC 28.1", "Twilio Webhook Receives Message", "Component", "WhatsApp"],
            ["37", "ISR-226", "TC 28.2", "Voice Note to Transcriber", "Integration", "WhatsApp"],
            ["38", "ISR-226", "TC 28.3", "WhatsApp Media Download", "Component", "WhatsApp"],
            ["39", "ISR-226", "TC 28.4", "Twilio Signature Validation", "Component", "WhatsApp"],
            ["40", "ISR-352", "TC 28.5", "Multi-Message Grouping", "Integration", "WhatsApp"],
            ["41", "ISR-227", "TC 29.1", "WhatsApp Auto-Acknowledgment", "Component", "WhatsApp"],
            ["42", "ISR-339", "TC 30.1", "Twilio Recording Webhook", "Component", "Phone Calls"],
            ["43", "ISR-340", "TC 30.3", "Recording Failure Doesn't Block Call", "Component", "Phone Calls"],
            ["44", "ISR-365", "TC 36.1", "Name List AI Recognition", "Component", "AI Output"],
            ["45", "ISR-409", "TC 36.5", "Gemma Model Performance", "Component", "AI Output"],
            ["46", "ISR-375", "TC 36.6", "Profanity Filtering", "Component", "AI Output"],
            ["47", "ISR-232", "TC 34.1", "Word Doc Repak Format", "Component", "Word Generator"],
            ["48", "ISR-232", "TC 34.2", "Bulk Word Generation", "Integration", "Word Generator"],
            ["49", "ISR-401", "TC 37.1", "Login Gates All Routes", "Component", "Dashboard"],
            ["50", "ISR-401", "TC 37.2", "Invalid Credentials Rejected", "Component", "Dashboard"],
            ["51", "ISR-401", "TC 37.4", "Admin Account Management", "Acceptance", "Dashboard"],
            ["52", "ISR-400", "TC 38.1", "Incident List Filtering", "Component", "Dashboard"],
            ["53", "ISR-400", "TC 38.3", "Semantic Search", "Integration", "Dashboard"],
        ],
    )

    # ── 5. Additional Risks ──────────────────────────────────────────────
    doc.add_heading("Addendum: Additional Risks (Sprint 5-6)", level=1)

    add_table(
        doc,
        ["Risk", "Probability", "Impact", "Mitigation"],
        [
            [
                "R5: Email classifier false positives create junk cases",
                "3",
                "3",
                "Test with 30+ real emails, tune threshold, add engineer override",
            ],
            [
                "R6: WhatsApp Twilio signature validation fails in production",
                "2",
                "4",
                "Test with real Twilio account, document configuration",
            ],
            [
                "R7: Gemma model quality regression vs llama3.1",
                "3",
                "3",
                "Benchmark both models with same test data before switching",
            ],
            [
                "R8: Dashboard session bypass exposes customer data",
                "2",
                "5",
                "Security review of middleware, test with browser dev tools",
            ],
        ],
    )

    # ── Save ─────────────────────────────────────────────────────────────
    doc.save(OUTPUT_PATH)
    print(f"Done. Saved to {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
