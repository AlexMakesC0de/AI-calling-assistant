from docx import Document

doc = Document(r"D:\Downloads\Test Plan v1.1.docx")

print("=== HEADINGS ===")
for p in doc.paragraphs:
    if p.style.name.startswith("Heading"):
        print(p.style.name + ": " + p.text[:120])

print()
print("Total tables: " + str(len(doc.tables)))
print()
print("=== TABLE STRUCTURES ===")
for t_idx, table in enumerate(doc.tables):
    rows = len(table.rows)
    cols = len(table.columns)
    print("--- Table " + str(t_idx + 1) + " (" + str(rows) + " rows x " + str(cols) + " cols) ---")
    headers = []
    for cell in table.rows[0].cells:
        headers.append(cell.text.strip()[:50])
    print("  Headers: " + " | ".join(headers))
    if rows >= 2:
        r2 = []
        for cell in table.rows[1].cells:
            r2.append(cell.text.strip()[:50])
        print("  Row 2:   " + " | ".join(r2))
    if rows >= 3:
        r3 = []
        for cell in table.rows[2].cells:
            r3.append(cell.text.strip()[:50])
        print("  Row 3:   " + " | ".join(r3))
