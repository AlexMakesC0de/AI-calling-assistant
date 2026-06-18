from docx import Document

doc = Document(r"D:\Downloads\Test Plan v1.1.docx")

# Read full content of Table 8 (test case index) - all rows
print("=== TABLE 8: TEST CASE INDEX (all rows) ===")
t8 = doc.tables[7]
for r_idx, row in enumerate(t8.rows):
    cells = []
    for cell in row.cells:
        cells.append(cell.text.strip()[:60])
    print("  Row " + str(r_idx) + ": " + " | ".join(cells))

print()

# Read full content of Table 9 (first test case spec card)
print("=== TABLE 9: FIRST TEST CASE SPEC ===")
t9 = doc.tables[8]
for r_idx, row in enumerate(t9.rows):
    cells = []
    for cell in row.cells:
        cells.append(cell.text.strip()[:80])
    print("  Row " + str(r_idx) + ": " + " | ".join(cells))

print()

# Read Table 33 (last test case spec - Graph API)
print("=== TABLE 33: GRAPH API TEST CASE SPEC ===")
t33 = doc.tables[32]
for r_idx, row in enumerate(t33.rows):
    cells = []
    for cell in row.cells:
        cells.append(cell.text.strip()[:80])
    print("  Row " + str(r_idx) + ": " + " | ".join(cells))

print()

# Read Table 34 and 35 too
print("=== TABLE 34 ===")
t34 = doc.tables[33]
for r_idx, row in enumerate(t34.rows):
    cells = []
    for cell in row.cells:
        cells.append(cell.text.strip()[:80])
    print("  Row " + str(r_idx) + ": " + " | ".join(cells))

print()
print("=== TABLE 35 ===")
t35 = doc.tables[34]
for r_idx, row in enumerate(t35.rows):
    cells = []
    for cell in row.cells:
        cells.append(cell.text.strip()[:80])
    print("  Row " + str(r_idx) + ": " + " | ".join(cells))

# Also check Version history table (Table 1)
print()
print("=== TABLE 1: VERSION HISTORY (all rows) ===")
t1 = doc.tables[0]
for r_idx, row in enumerate(t1.rows):
    cells = []
    for cell in row.cells:
        cells.append(cell.text.strip()[:60])
    print("  Row " + str(r_idx) + ": " + " | ".join(cells))

# Table 2: Test items
print()
print("=== TABLE 2: TEST ITEMS (all rows) ===")
t2 = doc.tables[1]
for r_idx, row in enumerate(t2.rows):
    cells = []
    for cell in row.cells:
        cells.append(cell.text.strip()[:60])
    print("  Row " + str(r_idx) + ": " + " | ".join(cells))

# Table 3: Requirements (functional)
print()
print("=== TABLE 3: FUNCTIONAL REQUIREMENTS (all rows) ===")
t3 = doc.tables[2]
for r_idx, row in enumerate(t3.rows):
    cells = []
    for cell in row.cells:
        cells.append(cell.text.strip()[:60])
    print("  Row " + str(r_idx) + ": " + " | ".join(cells))

# Table 5: Acceptance criteria
print()
print("=== TABLE 5: ACCEPTANCE CRITERIA (all rows) ===")
t5 = doc.tables[4]
for r_idx, row in enumerate(t5.rows):
    cells = []
    for cell in row.cells:
        cells.append(cell.text.strip()[:80])
    print("  Row " + str(r_idx) + ": " + " | ".join(cells))

# Risks table 39
print()
print("=== TABLE 39: KEY RISKS ===")
t39 = doc.tables[38]
for r_idx, row in enumerate(t39.rows):
    cells = []
    for cell in row.cells:
        cells.append(cell.text.strip()[:60])
    print("  Row " + str(r_idx) + ": " + " | ".join(cells))
