"""
modify_test_report_xlsx.py

Opens the existing Test Report at D:/Downloads/Test Report.xlsx,
appends Sprint 5-6 test case rows after existing data, applies styling
(gray "Not Run" status, thin borders, alternating light-blue/white rows),
and saves in place.
"""

import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
INPUT_PATH = r"D:\Downloads\Test Report.xlsx"
SHEET_NAME = "Test Cases"

# Styling constants
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
GRAY_FILL = PatternFill(start_color="C0C0C0", end_color="C0C0C0", fill_type="solid")
LIGHT_BLUE_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
WRAP_ALIGNMENT = Alignment(wrap_text=True, vertical="top")

# ---------------------------------------------------------------------------
# New test case data
# Columns: CaseNumber, TestScenario, Description, TestItem, TestLevel,
#           TestTechniques, Tester, Steps, Expected Results, Status
# ---------------------------------------------------------------------------
NEW_ROWS = [
    # --- Outlook Email Pipeline (ISR-216) ---
    [
        "Sprint5#28",
        "Email Pipeline E2E",
        "Given inbound Outlook email... service form produced",
        "Email Monitor, Voice App, Formatter",
        "Integration",
        "Scenario-based, Workflow",
        "",
        "Send test email to configured inbox, verify detection, pipeline trigger, service form creation",
        'Email detected, pipeline triggered, service form created with source_channel="email"',
        "Not Run",
    ],
    [
        "Sprint5#29",
        "Email Body to Formatter (bypass transcriber)",
        "Given email text input... transcriber skipped",
        "Voice App, Formatter",
        "Component",
        "Workflow, Data-focused",
        "",
        "Submit email body to /ingest/email, verify transcriber not called, formatter receives text",
        "Formatter processes email text directly, transcriber logs show no activity",
        "Not Run",
    ],
    [
        "Sprint5#30",
        "Email Classification",
        "Given inbound emails of different types... classified correctly",
        "Email Monitor, Classifier",
        "Component",
        "Data-focused, Scenario",
        "",
        "Send support email, newsletter, spam. Check classifier output",
        "Support=support, newsletter=not_support, ambiguous=unclear with confidence score",
        "Not Run",
    ],
    [
        "Sprint5#31",
        "Low-Confidence Email Flagging",
        "Given unclear email... flagged for review",
        "Classifier, Dashboard",
        "Integration",
        "Data-focused, Workflow",
        "",
        "Send ambiguous email, verify low_confidence flag set, visible in dashboard",
        "Low-confidence flag set when score < threshold, visible for engineer review",
        "Not Run",
    ],
    [
        "Sprint5#32",
        "Email Attachments Stored Per Case",
        "Given email with attachments... downloaded and stored",
        "Email Monitor, Graph API",
        "Component",
        "Data-focused, Scenario",
        "",
        "Send email with image and PDF attachment, check storage",
        "Attachments downloaded, stored per case, referenced in form",
        "Not Run",
    ],
    [
        "Sprint5#33",
        "Oversized Attachment Handling",
        "Given email with large attachment... logged and notified",
        "Email Monitor",
        "Integration",
        "Data-focused, Config",
        "",
        "Send email with >10MB attachment, check logs",
        "Oversized attachment logged, engineer notified, not silently dropped",
        "Not Run",
    ],
    [
        "Sprint5#34",
        "Email Signature Stripping",
        "Given email with signature/quoted replies... stripped before AI",
        "Email Monitor",
        "Component",
        "Data-focused",
        "",
        "Send reply email with quoted text and signature, check formatter input",
        "Only new content sent to AI, signatures and quoted replies removed",
        "Not Run",
    ],
    [
        "Sprint5#35",
        "Plain-text and MIME Email Support",
        "Given emails in different formats... all processed",
        "Email Monitor",
        "Component",
        "Data-focused, Config",
        "",
        "Send plain-text, HTML, multipart MIME emails, verify all processed",
        "All email formats correctly parsed and processed",
        "Not Run",
    ],
    # --- WhatsApp (ISR-217) ---
    [
        "Sprint5#36",
        "Twilio Webhook Receives Message",
        "Given WhatsApp text message... stored in DB",
        "WhatsApp Ingest, Database",
        "Component",
        "Scenario, Data-focused",
        "",
        "Send WhatsApp message via Twilio, check DB",
        "Message stored in WhatsAppMessage table with correct fields",
        "Not Run",
    ],
    [
        "Sprint5#37",
        "Voice Note to Transcriber",
        "Given WhatsApp voice note... forwarded to transcriber",
        "WhatsApp Ingest, Transcriber",
        "Integration",
        "Workflow, Data-focused",
        "",
        "Send voice note, monitor transcriber logs",
        "Audio forwarded to transcriber, transcript produced",
        "Not Run",
    ],
    [
        "Sprint5#38",
        "WhatsApp Media Download",
        "Given WhatsApp image/PDF... downloaded and stored",
        "WhatsApp Ingest",
        "Component",
        "Data-focused, Scenario",
        "",
        "Send image and PDF via WhatsApp, check local storage",
        "Media downloaded from Twilio URL, stored locally",
        "Not Run",
    ],
    [
        "Sprint5#39",
        "Twilio Signature Validation",
        "Given forged webhook request... rejected",
        "WhatsApp Ingest",
        "Component",
        "Config, Scenario",
        "",
        "Send request without valid Twilio signature",
        "Request rejected with 403, logged",
        "Not Run",
    ],
    [
        "Sprint5#40",
        "Multi-Message Grouping",
        "Given multiple messages in 10-min window... grouped",
        "WhatsApp Ingest",
        "Integration",
        "Workflow, Data-focused",
        "",
        "Send 3 messages within 10 min, check grouping",
        "Messages grouped into single conversation thread",
        "Not Run",
    ],
    [
        "Sprint5#41",
        "WhatsApp Auto-Acknowledgment",
        "Given first message from contact... auto-reply sent",
        "WhatsApp Ingest",
        "Component",
        "Scenario, Config",
        "",
        "Send first message from new number, check auto-reply",
        "Auto-acknowledgment sent, text matches WHATSAPP_AUTO_ACK_TEXT",
        "Not Run",
    ],
    # --- Phone Calls (ISR-218) ---
    [
        "Sprint6#42",
        "Twilio Recording Webhook",
        "Given Twilio call recording... forwarded to pipeline",
        "Telephony Ingest",
        "Component",
        "Scenario, Config",
        "",
        "Trigger Twilio recording webhook, check forwarding",
        "Recording forwarded to voice-app /upload endpoint",
        "Not Run",
    ],
    [
        "Sprint6#43",
        "Recording Failure Doesn't Block Call",
        "Given recording fails... call continues",
        "Telephony Ingest",
        "Component",
        "Workflow, Config",
        "",
        "Simulate recording failure during call",
        "Call completes normally, failure logged",
        "Not Run",
    ],
    # --- AI Output Reliability (ISR-221) ---
    [
        "Sprint6#44",
        "Name List Improves AI Recognition",
        "Given name list configured... names correctly filled",
        "Formatter, Name List",
        "Component",
        "Data-focused, Scenario",
        "",
        "Configure name list, process transcript with known names",
        "Names correctly recognized and filled in form",
        "Not Run",
    ],
    [
        "Sprint6#45",
        "Gemma Model Performance Test",
        "Given Gemma model vs llama3.1... benchmarked",
        "Formatter, Ollama",
        "Component",
        "Data-focused, Config",
        "",
        "Process same transcripts with both models, compare output quality",
        "Performance comparison documented with accuracy metrics",
        "Not Run",
    ],
    [
        "Sprint6#46",
        "Profanity Filtering",
        "Given transcript with profanity... filtered",
        "Formatter, Content Filter",
        "Component",
        "Data-focused, Scenario",
        "",
        "Process transcript containing profanity, check output",
        "Profanity replaced with [redacted] per TERM_FILTER_MODE",
        "Not Run",
    ],
    # --- Word Generator (ISR-209) ---
    [
        "Sprint6#47",
        "Word Doc Problem/Analysis/Advice Format",
        "Given incident form... Word generated in Repak format",
        "Word Generator",
        "Component",
        "Data-focused, Scenario",
        "",
        "Generate Word doc from incident, check structure",
        "Document has Problem/Analysis/Advice sections matching Repak template",
        "Not Run",
    ],
    [
        "Sprint6#48",
        "Bulk Word Generation",
        "Given multiple incidents... batch Word doc produced",
        "Word Generator, Dashboard",
        "Integration",
        "Workflow, Scenario",
        "",
        "Select date range, generate bulk Word document",
        "Multi-incident document generated with all selected incidents",
        "Not Run",
    ],
    # --- Dashboard (ISR-222) ---
    [
        "Sprint6#49",
        "Login Gates All Routes",
        "Given unauthenticated user... redirected to login",
        "Dashboard, Middleware",
        "Component",
        "Config, Scenario",
        "",
        "Access /incidents without login, check redirect",
        "Redirected to /login, no data exposed",
        "Not Run",
    ],
    [
        "Sprint6#50",
        "Invalid Credentials Rejected",
        "Given wrong password... error shown",
        "Dashboard, Auth",
        "Component",
        "Scenario, Data-focused",
        "",
        "Enter wrong credentials on login page",
        "Clear error message, no access granted",
        "Not Run",
    ],
    [
        "Sprint6#51",
        "Admin Create/Delete Accounts",
        "Given admin user... can manage accounts",
        "Dashboard, Admin Panel",
        "Acceptance",
        "Scenario, Usability",
        "",
        "Create user, verify login works, delete user",
        "User created, can login, deleted successfully",
        "Not Run",
    ],
    [
        "Sprint6#52",
        "Incident List Filtering",
        "Given incidents page... filterable by category/priority",
        "Dashboard",
        "Component",
        "Usability, Scenario",
        "",
        "Apply category and priority filters on /incidents",
        "List updates to show only matching incidents",
        "Not Run",
    ],
    [
        "Sprint6#53",
        "Semantic Search",
        "Given search query... relevant results returned",
        "Dashboard, pgvector",
        "Integration",
        "Data-focused, Workflow",
        "",
        "Search for a known term, check results",
        "Relevant incidents returned ranked by cosine similarity",
        "Not Run",
    ],
]


def main() -> None:
    wb = openpyxl.load_workbook(INPUT_PATH)
    ws = wb[SHEET_NAME]

    start_row = ws.max_row + 1  # first empty row after existing data
    num_cols = 10  # A-J

    print(f"Existing data ends at row {ws.max_row}. Appending {len(NEW_ROWS)} rows starting at row {start_row}.")

    for idx, row_data in enumerate(NEW_ROWS):
        excel_row = start_row + idx

        # Determine alternating fill: even idx = light blue, odd idx = white
        row_fill = LIGHT_BLUE_FILL if idx % 2 == 0 else WHITE_FILL

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=excel_row, column=col_idx, value=value)
            cell.border = THIN_BORDER
            cell.alignment = WRAP_ALIGNMENT

            # Status column (col 10) gets gray fill for "Not Run"
            if col_idx == 10:
                cell.fill = GRAY_FILL
            else:
                cell.fill = row_fill

    # Save
    wb.save(INPUT_PATH)
    print(f"Saved {len(NEW_ROWS)} new rows to '{INPUT_PATH}'.")
    print(f"Data now spans rows 2-{start_row + len(NEW_ROWS) - 1} (row 1 = headers).")


if __name__ == "__main__":
    main()
