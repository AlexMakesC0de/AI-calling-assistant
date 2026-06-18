from docx import Document

doc = Document(r"D:\Downloads\Test Plan v1.1.docx")

# Table 39 (index 38)
t = doc.tables[38]
print("Table 39 - columns: " + str(len(t.columns)) + ", rows: " + str(len(t.rows)))
row = t.add_row()
print("New row cells count: " + str(len(row.cells)))

# Try setting cells
cells_text = ["R5", "Test risk", "3", "3", "Mitigation text"]
for i, text in enumerate(cells_text):
    print("Setting cell " + str(i) + " of " + str(len(row.cells)))
    cell = row.cells[i]
    cell.text = text
print("Table 39 OK")

# Table 40 (index 39)
t40 = doc.tables[39]
print("Table 40 - columns: " + str(len(t40.columns)) + ", rows: " + str(len(t40.rows)))
row40 = t40.add_row()
print("New row cells count: " + str(len(row40.cells)))

cells_text40 = ["Risk text", "Impact text", "Mitigation text"]
for i, text in enumerate(cells_text40):
    print("Setting cell " + str(i) + " of " + str(len(row40.cells)))
    cell = row40.cells[i]
    cell.text = text
print("Table 40 OK")
